"""Build anonymous Gate 5 project brief and technical report DOCX files."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "deliverables"

INK = "172B4D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6573"
LIGHT = "F2F4F7"
BORDER = "C9D1DB"
WHITE = "FFFFFF"
TOTAL_DXA = 9360


def _set_run_font(run, size: float, *, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int], *, indent: int = 120) -> None:
    if sum(widths) != TOTAL_DXA:
        raise ValueError(f"table widths must total {TOTAL_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(TOTAL_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_cell_text(cell, text: str, *, bold: bool = False, color: str = "000000", center: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    _set_run_font(paragraph.add_run(text), 9.2, bold=bold, color=color)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _set_table_geometry(table, widths)
    _set_repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        _set_cell_shading(table.rows[0].cells[index], LIGHT)
        _set_cell_text(table.rows[0].cells[index], header, bold=True, color=INK, center=True)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            _set_cell_text(cells[index], value, center=index > 0 and len(value) < 22)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(paragraph.add_run("第 "), 8.5, color=MUTED)
    field_run = paragraph.add_run()
    _set_run_font(field_run, 8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.append(begin)
    field_run._r.append(instruction)
    field_run._r.append(end)
    _set_run_font(paragraph.add_run(" 页"), 8.5, color=MUTED)


def _configure_document(doc: Document, *, running_title: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    _set_run_font(header.add_run(running_title), 8.5, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    _add_page_number(footer)

    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.company = ""
    doc.core_properties.comments = "Anonymous competition submission"


def _add_title_block(doc: Document, title: str, subtitle: str, *, compact: bool = False) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(20 if compact else 72)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(12)
    _set_run_font(kicker.add_run("实时视觉智能 · 匿名提交"), 10, bold=True, color=BLUE)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    _set_run_font(paragraph.add_run(title), 25 if compact else 29, bold=True, color=INK)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(24 if compact else 42)
    _set_run_font(paragraph.add_run(subtitle), 12.5, color=MUTED)


def _add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.widow_control = True
    if bold_lead and text.startswith(bold_lead):
        _set_run_font(paragraph.add_run(bold_lead), 11, bold=True, color=INK)
        _set_run_font(paragraph.add_run(text[len(bold_lead) :]), 11)
    else:
        _set_run_font(paragraph.add_run(text), 11)


def _add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        _set_run_font(paragraph.add_run(item), 11)


def _add_metric_strip(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    _set_table_geometry(table, [2340, 2340, 2340, 2340])
    _set_repeat_header(table.rows[0])
    table.style = "Table Grid"
    values = [
        ("3.000M", "总参数"),
        ("18.28 ms", "V100 P95"),
        ("42.93%", "OF-Syn test MAP"),
        ("207", "自动化测试"),
    ]
    for cell, (value, label) in zip(table.rows[0].cells, values, strict=True):
        _set_cell_shading(cell, LIGHT)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        _set_run_font(p.add_run(value), 14, bold=True, color=INK)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        _set_run_font(p.add_run(label), 8.5, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_brief(path: Path) -> str:
    intro = (
        "本项目面向低算力平台构建纯视觉实时跌倒检测系统。系统以YOLO11n-pose提取人体骨架，"
        "结合轻量因果TCN建模动作时序，并通过多目标跟踪与事件聚合输出可追溯告警。模型共300.02万参数，"
        "FP32参数体积12.00MB；在Tesla V100、1920×1080输入下端到端P95为18.28ms。"
        "冻结模型在OF-Syn一次性测试集上的MAP为42.93%，P@R90/P@R95为45.66%/40.19%。"
        "方案支持无渲染批处理、断点恢复、哈希校验和匿名部署；已识别遮挡与跨域误报为主要局限。"
    )
    count = len(re.sub(r"\s+", "", intro))
    if count > 300:
        raise ValueError(f"brief exceeds 300 characters: {count}")

    doc = Document()
    _configure_document(doc, running_title="项目简介 | 匿名提交")
    _add_title_block(doc, "面向低算力平台的实时跌倒检测", "纯视觉 · 因果时序 · 可审计工程闭环", compact=True)
    _add_metric_strip(doc)
    heading = doc.add_paragraph("项目简介", style="Heading 1")
    heading.paragraph_format.space_before = Pt(20)
    _add_body(doc, intro)
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(14)
    note.paragraph_format.space_after = Pt(0)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(note.add_run(f"简介正文 {count} 字（含标点，去除空白）"), 9, color=MUTED)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return f"brief_chars={count}"


def build_report(path: Path) -> None:
    doc = Document()
    _configure_document(doc, running_title="实时跌倒检测系统 | 项目文档 | 匿名提交")
    _add_title_block(
        doc,
        "面向低算力端侧平台\n基于视觉的实时跌倒检测",
        "技术方案、评测证据与工程部署说明",
    )
    _add_metric_strip(doc)
    _add_body(
        doc,
        "提交版本：冻结 TCN-only epoch 7。所有测试数字均来自已归档产物；OF-Syn test 已完成唯一一次消费，后续不再调模型、阈值、融合或性能配置。",
    )
    doc.add_page_break()

    doc.add_paragraph("1. 项目概述", style="Heading 1")
    _add_body(
        doc,
        "本项目解决普通 RGB 或灰度视频中的实时跌倒事件检测。系统不依赖深度相机、穿戴设备或云端服务，"
        "以人体姿态时序为核心，在轻量参数预算内兼顾高召回、可解释告警与可复现批处理。评测重点对应准确性、模型体积与端到端时延三类指标。",
    )
    _add_bullets(
        doc,
        [
            "纯视觉输入：支持本地 MP4 与归档成员，推理过程可完全在本机或指定 GPU 服务器内闭环。",
            "事件级输出：每条告警包含 track_id、起止时间、峰值分数与模型来源，可供复核和下游联动。",
            "工程边界：缓存、断点恢复、签名、哈希和一次性 test seal 共同保证评测过程可审计。",
        ],
    )

    doc.add_paragraph("2. 系统架构", style="Heading 1")
    _add_table(
        doc,
        ["阶段", "主要模块", "输出与职责"],
        [
            ["1", "视频源与解码", "统一读取普通视频和 tar 成员；显式临时目录与内容哈希"],
            ["2", "YOLO11n-pose", "逐帧检测人体并估计 17 个 COCO 关键点"],
            ["3", "轻量多目标跟踪", "以 IoU 和中心运动维持 track_id，处理短时漏检"],
            ["4", "FallTCN", "对每条轨迹的因果骨架窗口输出跌倒概率"],
            ["5", "事件聚合", "平滑、双阈值、最小时长与间隔合并，生成告警事件"],
        ],
        [1100, 2600, 5660],
    )
    _add_body(
        doc,
        "在线路径只使用历史及当前帧，不读取未来信息。默认 TCN-only 模式避免了规则融合在遮挡压力下的不稳定；规则分支保留为诊断工具，不进入冻结提交候选。",
    )

    doc.add_paragraph("3. 模型与特征设计", style="Heading 1")
    _add_body(
        doc,
        "每帧将 17 个关键点表示为置信度加权、以躯干中心平移并按人体框宽高归一化的 (x, y, confidence)，形成 51 维特征。"
        "连续 16 帧组成因果窗口，轻量 TCN 通过扩张一维卷积提取动作变化并输出二分类 logit。该表示对画面平移和尺度变化较稳健，"
        "同时保持极低参数量。",
    )
    _add_body(
        doc,
        "结构探索严格采用单变量门控。fixed-time/global、难负与 MIL、多语义辅助头、joint+bone+motion+bbox 多流等候选均先做配对 pilot；"
        "没有达到预注册增益或统计稳定性门槛的方案立即停止。最终保留 Gate 3C epoch 7，是精度证据、鲁棒性和工程风险的联合决策。",
    )

    doc.add_paragraph("4. 数据与评测协议", style="Heading 1")
    _add_table(
        doc,
        ["数据", "用途", "规模", "协议"],
        [
            ["OF-Syn train", "训练", "9,600 clips", "仅训练；密集动作标签转窗口监督"],
            ["OF-Syn val", "选择与分析", "1,200 clips", "模型/阈值冻结、错误切片与扰动测试"],
            ["OF-Syn test", "最终一次性报告", "1,200 clips", "seal 约束；完成后永久禁止调参"],
            ["URFD val", "外部工程 fixture", "1 个冻结片段", "V100 1080P/300 帧端到端验收"],
        ],
        [2000, 1850, 1650, 3860],
    )
    _add_body(
        doc,
        "OF-Syn 使用上游固定 random clip split，因此不能宣称 subject-independent 或 scene-independent。"
        "本项目 MAP、P@R90 和 P@R95 均为一致的本地 clip-level 协议；官方 event matching 细则若不同，应在提交说明中明确口径差异。",
    )

    doc.add_paragraph("5. 冻结结果", style="Heading 1")
    _add_table(
        doc,
        ["数据/环境", "MAP", "P@R90", "P@R95", "备注"],
        [
            ["OF-Syn val", "49.99%", "52.77%", "47.20%", "阈值与候选冻结依据"],
            ["OF-Syn test", "42.93%", "45.66%", "40.19%", "1,200/1,200；唯一一次"],
            ["test - val", "-7.06点", "-7.11点", "-7.01点", "只报告，不回流调参"],
        ],
        [1900, 1350, 1350, 1350, 3410],
    )
    _add_body(
        doc,
        "test 的 R90 工作点为 TP/FP/FN=200/238/22，R95 工作点为 211/314/11。结果说明模型能维持高召回，"
        "但相似日常动作与跨划分差异会显著增加误报；这是当前性能上限的主要来源。",
    )

    doc.add_paragraph("6. 工程性能与合规", style="Heading 1")
    _add_table(
        doc,
        ["项目", "冻结结果", "约束", "结论"],
        [
            ["总参数", "3,000,165", "≤20M", "通过"],
            ["FP32 参数体积", "12.00066 MB", "≤80MB", "通过"],
            ["V100 端到端 P50/P95", "16.56/18.28 ms", "P95≤100ms", "通过"],
            ["输入与帧数", "1920×1080 / 300", "1080P 以上输入", "通过"],
            ["自动化质量门", "207 tests + Ruff", "全量回归", "通过"],
        ],
        [2350, 2700, 2450, 1860],
    )
    _add_body(
        doc,
        "V100 数据来自 Tesla V100-SXM2-16GB 的 PyTorch eager 正式复测，包含解码、检测、CPU transfer、跟踪、TCN 与事件聚合；"
        "不把本机 RTX 5070 Ti 结果冒充赛事硬件结果。模型权重、运行配置和报告均以 SHA-256 固定。",
    )

    doc.add_paragraph("7. 鲁棒性与失败模式", style="Heading 1")
    _add_table(
        doc,
        ["条件", "TCN 子集 MAP", "R95 Precision / Recall", "观察"],
        [
            ["clean", "39.71%", "39.0% / 94.1%", "固定 100-clip 扰动子集"],
            ["grayscale", "46.30%", "41.0% / 94.1%", "灰度输入未造成系统性回退"],
            ["lowlight", "34.17%", "36.4% / 94.1%", "低照度增加分数漂移"],
            ["Gaussian noise", "34.31%", "40.0% / 94.1%", "轻噪声主要影响精度"],
            ["center occlusion", "17.49%", "23.0% / 82.4%", "明确失效模式"],
        ],
        [2150, 1900, 2560, 2750],
    )
    _add_bullets(
        doc,
        [
            "主要假阳性来自 lie_down、lying、stand_up 等与跌倒形态相似的日常动作。",
            "中心遮挡同时破坏姿态检测、track 连续性和 TCN 输入，可能产生漏检与碎片化检测。",
            "当前 tracker 无 ReID；多人交叉、长时间遮挡或离场重入可能导致身份切换。",
            "现有红外证据主要是灰度代理，不等同于真实热红外专项验证。",
        ],
    )

    doc.add_paragraph("8. 与研究工作的关系", style="Heading 1")
    _add_body(
        doc,
        "ST-GCN、2s-AGCN 与 PoseC3D 分别证明了时空骨架关系、joint/bone 多流和置信度感知姿态表示的价值。"
        "本项目受这些方法启发，但没有直接复制较重图网络或 3D 卷积；在 20M 参数和 100ms 约束下，采用更小的因果 TCN，"
        "并以 pilot 门控验证 fixed-time、多语义和多流特征。",
    )
    _add_body(
        doc,
        "LFD-YOLO 与 BMR-YOLO 报告了单帧外观检测在低光、遮挡和相似动作上的改进，但其数据、标签和 mAP 定义与本项目不同，"
        "不能直接进行数值排名。本项目的差异化在于：以姿态时序替代单帧 fall/down 外观二分类，并给出完整 V100 端到端、"
        "高召回指标、扰动切片和一次性测试证据。",
    )

    doc.add_paragraph("9. 部署与复现", style="Heading 1")
    _add_bullets(
        doc,
        [
            "安装：Python 3.11 环境执行 pip install -r requirements.txt。",
            "单视频：infer.py 支持无渲染推理、事件 JSON 与可选可视化视频。",
            "批处理：eval.evaluate_manifest 支持签名化 resume、失败隔离和稳定 JSONL 输出。",
            "生产边界：部署时只允许冻结 checkpoint/run JSON；不得重新打开已完成的 OF-Syn test seal。",
            "隐私边界：视频、姿态与事件均可在本地进程处理；提交包不包含数据集或用户绝对路径。",
        ],
    )

    doc.add_paragraph("10. 后续改进方向", style="Heading 1")
    _add_body(
        doc,
        "测试封印后不再改变本次提交候选。未来独立版本应以新的 train/val/test 协议开展：优先补充真实红外、遮挡和跨场景数据；"
        "在 pose/ID 证据充分时比较 RTMO/RTMPose 或带 ReID 跟踪；使用固定时间采样和显式 observed mask；"
        "并在独立验证集上评估校准、难负挖掘与轻量骨架图网络。任何收益都应采用配对多种子和置信区间报告。",
    )

    doc.add_paragraph("参考文献", style="Heading 1")
    references = [
        "[1] Yan, Xiong, Lin. Spatial Temporal Graph Convolutional Networks for Skeleton-Based Action Recognition. AAAI 2018. https://arxiv.org/abs/1801.07455",
        "[2] Shi et al. Two-Stream Adaptive Graph Convolutional Networks for Skeleton-Based Action Recognition. CVPR 2019. https://openaccess.thecvf.com/content_CVPR_2019/html/Shi_Two-Stream_Adaptive_Graph_Convolutional_Networks_for_Skeleton-Based_Action_Recognition_CVPR_2019_paper.html",
        "[3] Duan et al. Revisiting Skeleton-Based Action Recognition. CVPR 2022. https://openaccess.thecvf.com/content/CVPR2022/html/Duan_Revisiting_Skeleton-Based_Action_Recognition_CVPR_2022_paper.html",
        "[4] OmniFall: A Unified Benchmark for Fall Detection and Activity Recognition. arXiv:2505.19889. https://arxiv.org/abs/2505.19889",
        "[5] LFD-YOLO: a lightweight fall detection model. Scientific Reports 15, 5069 (2025). https://doi.org/10.1038/s41598-025-89214-7",
        "[6] BMR-YOLO: fall detection in complex environments. PLOS ONE 20(11):e0335992 (2025). https://doi.org/10.1371/journal.pone.0335992",
    ]
    for reference in references:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.paragraph_format.space_after = Pt(5)
        _set_run_font(paragraph.add_run(reference), 9.2, color=MUTED)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    brief_path = OUTPUT_DIR / "fall_detection_project_brief.docx"
    report_path = OUTPUT_DIR / "fall_detection_project_report.docx"
    brief_result = build_brief(brief_path)
    build_report(report_path)
    print(brief_result)
    print(brief_path)
    print(report_path)


if __name__ == "__main__":
    main()
